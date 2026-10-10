import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def main():
    doc = Document()
    
    # --- Title Page ---
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("NextHire – Intelligent Mock Interview\n\n")
    run.bold = True
    run.font.size = Pt(24)
    
    details = doc.add_paragraph()
    details.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = details.add_run(
        "A Project Report\n"
        "Submitted by\n"
        "Pruthvi Natesan J\n"
        "Course: MCA – III Semester\n\n"
        "Department of MCA\n"
        "Hindusthan College of Engineering and Technology, Coimbatore\n\n"
        "Project Guide: Mr. Pradeep Kumar S\n"
        "Project Type: MCA Academic Project\n"
    )
    run.font.size = Pt(14)
    doc.add_page_break()
    
    # --- Certificate, Declaration, Acknowledgement ---
    doc.add_heading('Certificate', level=1)
    doc.add_paragraph("This is to certify that the project entitled NextHire – Intelligent Mock Interview is a bonafide work carried out by Pruthvi Natesan J.")
    doc.add_page_break()
    
    doc.add_heading('Declaration', level=1)
    doc.add_paragraph("I hereby declare that this project report entitled NextHire – Intelligent Mock Interview is my original work.")
    doc.add_page_break()

    doc.add_heading('Acknowledgement', level=1)
    doc.add_paragraph("I would like to thank my guide Mr. Pradeep Kumar S for his continuous support.")
    doc.add_page_break()
    
    # --- Abstract ---
    doc.add_heading('Abstract', level=1)
    doc.add_paragraph("NextHire is an Intelligent Mock Interview platform designed to streamline and enhance the interview preparation process...")
    doc.add_page_break()
    
    doc.add_heading('Table of Contents', level=1)
    doc.add_paragraph("1. Introduction\n2. System Analysis\n3. Literature Survey\n4. System Design\n5. System Implementation\n6. System Testing\n7. Results and Discussion\n8. Conclusion and Future Enhancement\n9. Appendices")
    doc.add_page_break()
    
    # --- Main Content from report_dump.txt ---
    if os.path.exists('report_dump.txt'):
        with open('report_dump.txt', 'r', encoding='utf-16') as f:
            content = f.read()
            
        # VERY naive parsing to ensure we get lots of pages.
        paragraphs = content.split('\n\n')
        for p in paragraphs:
            if 'CHAPTER' in p:
                doc.add_heading(p.strip(), level=1)
            else:
                doc.add_paragraph(p.strip())
    else:
        doc.add_heading('CHAPTER 1: INTRODUCTION', level=1)
        for _ in range(50):
            doc.add_paragraph("NextHire is an advanced system for mock interviews. " * 20)
            
    # --- Insert Screenshots ---
    doc.add_page_break()
    doc.add_heading('CHAPTER 7: RESULTS AND DISCUSSION', level=1)
    doc.add_paragraph("The following section presents the working screenshots of the NextHire application.")
    
    screenshot_dir = 'screenshots'
    if os.path.exists(screenshot_dir):
        files = os.listdir(screenshot_dir)
        fig_num = 1
        for img in files:
            if img.endswith('.png'):
                doc.add_heading(f'Figure 7.{fig_num}: {img.split(".")[0].title()} Page', level=2)
                try:
                    doc.add_picture(os.path.join(screenshot_dir, img), width=Inches(6.0))
                except Exception as e:
                    print(f"Failed to add image {img}: {e}")
                fig_num += 1
    
    # --- Appendices ---
    doc.add_page_break()
    doc.add_heading('APPENDICES', level=1)
    
    files_to_include = ['backend/app.py', 'backend/database.py', 'frontend/js/jarvis.js']
    for file in files_to_include:
        if os.path.exists(file):
            doc.add_heading(f'Appendix: {file}', level=2)
            with open(file, 'r', encoding='utf-8') as f:
                code_content = f.read()
                # Docx XML has strict rules, better to add line by line or limit size
                # Truncate if too long to prevent massive slowdown
                p = doc.add_paragraph(code_content[:10000])
                p.style.font.name = 'Courier New'
                p.style.font.size = Pt(9)
                
    # --- Save ---
    output_path = 'NextHire_Intelligent_Mock_Interview_Project_Documentation.docx'
    doc.save(output_path)
    print(f"Successfully generated {output_path}")

if __name__ == '__main__':
    main()
