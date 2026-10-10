"""
NextHire - Resume Parser Module
Extracts plain text and technical skills from PDF and DOCX files.
"""
import os
import re
from pathlib import Path
import fitz  # PyMuPDF
import docx  # python-docx

# Comprehensive Dictionary of Technical Skills for Mock Interview Evaluation
TECHNICAL_SKILLS = {
    "languages": [
        "python", "java", "javascript", "typescript", "c++", "c#", "c",
        "go", "golang", "rust", "php", "ruby", "swift", "kotlin", "scala",
        "sql", "html", "html5", "css", "css3", "bash", "shell"
    ],
    "frameworks": [
        "react", "react.js", "angular", "vue", "vue.js", "node.js", "nodejs",
        "express", "express.js", "flask", "django", "fastapi", "spring",
        "spring boot", "asp.net", ".net core", "next.js", "nuxt.js", "laravel"
    ],
    "databases": [
        "mysql", "postgresql", "postgres", "mongodb", "redis", "sqlite",
        "oracle", "sql server", "cassandra", "dynamodb", "elasticsearch"
    ],
    "devops_cloud": [
        "git", "github", "gitlab", "docker", "kubernetes", "aws", "amazon web services",
        "azure", "gcp", "google cloud", "ci/cd", "jenkins", "terraform", "linux", "nginx"
    ],
    "concepts": [
        "rest api", "restful", "microservices", "oop", "object oriented",
        "data structures", "algorithms", "system design", "graphql", "mvc", "agile", "scrum"
    ]
}

def extract_text_from_file(file_path: Path) -> str:
    """
    Extracts raw text content from either a PDF or DOCX file.
    """
    path_str = str(file_path).lower()
    text = ""

    if path_str.endswith(".pdf"):
        try:
            doc = fitz.open(file_path)
            for page in doc:
                text += page.get_text() + "\n"
            doc.close()
        except Exception:
            with open(file_path, "r", errors="ignore") as f:
                text = f.read()

    elif path_str.endswith(".docx") or path_str.endswith(".doc"):
        try:
            doc = docx.Document(file_path)
            for para in doc.paragraphs:
                text += para.text + "\n"
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + " "
                    text += "\n"
        except Exception as e:
            # Fallback for older .doc or plain text
            with open(file_path, "r", errors="ignore") as f:
                text = f.read()

    return text.strip()

def extract_skills_from_text(text: str) -> dict:
    """
    Scans the extracted text for technical skills across categories.
    """
    lower_text = " " + re.sub(r"[^\w\s\+\#\.\-]", " ", text.lower()) + " "
    detected = {
        "languages": [],
        "frameworks": [],
        "databases": [],
        "devops_cloud": [],
        "concepts": [],
        "all_skills": []
    }

    for category, skill_list in TECHNICAL_SKILLS.items():
        for skill in skill_list:
            # Use regex pattern to avoid partial word substring matches (e.g., 'c' inside 'cloud')
            escaped = re.escape(skill)
            pattern = rf"(?<!\w){escaped}(?!\w)"
            if re.search(pattern, lower_text):
                formatted = skill.title() if len(skill) > 3 else skill.upper()
                if skill in ["javascript", "typescript", "mysql", "mongodb", "postgresql", "fastapi", "docker", "kubernetes"]:
                    formatted = skill.title()
                elif skill in ["sql", "html", "css", "aws", "gcp", "ci/cd", "oop", "mvc"]:
                    formatted = skill.upper()

                if formatted not in detected[category]:
                    detected[category].append(formatted)
                    detected["all_skills"].append(formatted)

    return detected
